"""The household plan (the Homestead): monthly income, bills, savings and balances, kept in the UIP.

It replaces the "Income & Net Worth Tracker" workbook. A plan is one JSON document; every save is
kept as a new immutable version, and the newest version is the plan. Balances entered for a month
are anchors (actuals or the workbook's own starting points); months without one roll forward:

    bank        = previous bank + income + expenses (negative) - investment - house-fund contributions
                  (or, with a bank check-in for the month: the balance you entered + everything not yet
                  received or paid that month; an entered month-end bank balance still wins)
    retirement  = previous retirement x (1 + plan_retirement_return / 12) + retirement contribution
    investments = previous investments x (1 + plan_investment_return / 12) + investment contribution
    house fund  = previous house fund + house-fund contribution

This reproduces the workbook's Monthly Tracker exactly, which the importer checks.
"""
from __future__ import annotations

from contextlib import closing
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import io
import json
import re
import sqlite3
from typing import Callable

SCHEMA = "household-plan-1"
MONTH = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
BALANCES = ("bank_balance", "retirement_balance", "investment_balance", "house_fund_balance", "home_equity")
DEFAULT_SETTINGS = {
    "target_month": "2029-12",
    "home_price_low": 350000,
    "home_price_high": 450000,
    "down_payment_pct": 0.20,
    "closing_cost_pct": 0.03,
    "safety_fund": None,            # None: safety_fund_months of the target month's planned expenses
    "safety_fund_months": 6,
    "plan_investment_return": 0.10,
    "plan_retirement_return": 0.10,
    "contribution_mix": {"stocks": 1.0},
    "house_counts_retirement": False,
    "mortgage_rate": 0.065,         # placeholder from the workbook's House Purchase Planner, not a live rate
    "loan_years": 30,
    "home_price_growth": 0.0,       # yearly; 0 = today's prices. The Homestead offers the housing forecast.
}
MIX_SLEEVES = ("stocks", "crypto", "metals", "mtg", "cash")


class HouseholdPlanError(ValueError):
    pass


def _number(value, field, *, allow_none=False):
    if value is None or value == "":
        if allow_none:
            return None
        return 0.0
    try:
        n = float(value)
    except (TypeError, ValueError):
        raise HouseholdPlanError(f"{field} must be a number") from None
    if n != n or n in (float("inf"), float("-inf")):
        raise HouseholdPlanError(f"{field} must be a finite number")
    return round(n, 2)


def normalize_plan(document) -> dict:
    """Validate a plan document and return its canonical form (raises HouseholdPlanError)."""
    if not isinstance(document, dict):
        raise HouseholdPlanError("plan must be a JSON object")
    settings = dict(DEFAULT_SETTINGS)
    raw_settings = document.get("settings") or {}
    if not isinstance(raw_settings, dict):
        raise HouseholdPlanError("settings must be an object")
    for key in DEFAULT_SETTINGS:
        if key in raw_settings:
            settings[key] = raw_settings[key]
    if not MONTH.match(str(settings["target_month"])):
        raise HouseholdPlanError("target_month must look like 2029-12")
    for key in ("home_price_low", "home_price_high", "safety_fund_months"):
        settings[key] = _number(settings[key], key)
        if settings[key] < 0:
            raise HouseholdPlanError(f"{key} must not be negative")
    if settings["home_price_high"] < settings["home_price_low"]:
        raise HouseholdPlanError("home_price_high must be at least home_price_low")
    settings["safety_fund"] = _number(settings["safety_fund"], "safety_fund", allow_none=True)
    settings["loan_years"] = int(_number(settings["loan_years"], "loan_years"))
    if not 5 <= settings["loan_years"] <= 40:
        raise HouseholdPlanError("loan_years must be between 5 and 40")
    try:
        settings["home_price_growth"] = float(settings["home_price_growth"] or 0)
    except (TypeError, ValueError):
        raise HouseholdPlanError("home_price_growth must be a number") from None
    if not -0.2 <= settings["home_price_growth"] <= 0.3:
        raise HouseholdPlanError("home_price_growth is out of range")
    for key in ("down_payment_pct", "closing_cost_pct", "plan_investment_return", "plan_retirement_return", "mortgage_rate"):
        try:
            settings[key] = float(settings[key])
        except (TypeError, ValueError):
            raise HouseholdPlanError(f"{key} must be a number") from None
        if not -0.5 <= settings[key] <= 1:
            raise HouseholdPlanError(f"{key} is out of range")
    mix = settings["contribution_mix"]
    if not isinstance(mix, dict) or not mix:
        raise HouseholdPlanError("contribution_mix must name at least one sleeve")
    unknown = set(mix) - set(MIX_SLEEVES)
    if unknown:
        raise HouseholdPlanError(f"unknown contribution sleeve: {sorted(unknown)[0]}")
    weights = {k: float(v) for k, v in mix.items() if float(v) > 0}
    if any(v < 0 for v in weights.values()) or abs(sum(weights.values()) - 1) > 0.001:
        raise HouseholdPlanError("contribution_mix shares must add up to 100%")
    settings["contribution_mix"] = weights
    settings["house_counts_retirement"] = bool(settings["house_counts_retirement"])

    income_lines = [str(x).strip() for x in document.get("income_lines") or [] if str(x).strip()]
    expense_lines = [str(x).strip() for x in document.get("expense_lines") or [] if str(x).strip()]
    if len(set(income_lines)) != len(income_lines) or len(set(expense_lines)) != len(expense_lines):
        raise HouseholdPlanError("line names must be unique")
    months, seen = [], set()
    for raw in document.get("months") or []:
        if not isinstance(raw, dict):
            raise HouseholdPlanError("each month must be an object")
        month = str(raw.get("month", ""))
        if not MONTH.match(month):
            raise HouseholdPlanError(f"bad month: {month!r}")
        if month in seen:
            raise HouseholdPlanError(f"month {month} appears twice")
        seen.add(month)
        income = {line: _number((raw.get("income") or {}).get(line), f"{month} {line}") for line in income_lines}
        expenses = {line: _number((raw.get("expenses") or {}).get(line), f"{month} {line}") for line in expense_lines}
        item = {"month": month, "income": income, "expenses": expenses}
        for key in ("investment_contribution", "house_fund_contribution", "retirement_contribution"):
            item[key] = _number(raw.get(key), f"{month} {key}")
        for key in BALANCES:
            item[key] = _number(raw.get(key), f"{month} {key}", allow_none=True)
        item["note"] = str(raw.get("note") or "")[:500]
        item["bank_check"] = _bank_check(raw.get("bank_check"), month, income_lines, expense_lines)
        months.append(item)
    if not months:
        raise HouseholdPlanError("the plan has no months")
    months.sort(key=lambda m: m["month"])
    source = document.get("source") if isinstance(document.get("source"), dict) else {}
    return {"schema_version": SCHEMA, "settings": settings, "income_lines": income_lines,
            "expense_lines": expense_lines, "months": months, "source": source}


DATE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])-(0[1-9]|[12]\d|3[01])$")
TRANSFERS = ("investment_contribution", "house_fund_contribution")


def _bank_check(raw, month: str, income_lines: list[str], expense_lines: list[str]) -> dict | None:
    """A mid-month bank check-in: today's real balance and which of the month's lines already happened."""
    if raw in (None, "", {}):
        return None
    if not isinstance(raw, dict):
        raise HouseholdPlanError(f"{month} bank_check must be an object")
    balance = _number(raw.get("balance"), f"{month} bank check balance", allow_none=True)
    if balance is None:
        raise HouseholdPlanError(f"{month} bank check needs a balance")
    as_of = str(raw.get("as_of") or "")
    if not DATE.match(as_of) or as_of[:7] != month:
        raise HouseholdPlanError(f"{month} bank check date must be a day in that month")
    known = {f"income:{x}" for x in income_lines} | {f"expenses:{x}" for x in expense_lines} | set(TRANSFERS)
    done = sorted({str(x) for x in raw.get("done") or []})
    unknown = [x for x in done if x not in known]
    if unknown:
        raise HouseholdPlanError(f"{month} bank check names an unknown line: {unknown[0]}")
    return {"balance": balance, "as_of": as_of, "done": done}


def still_to_come(m: dict) -> float:
    """What a month's lines not yet done will add to the bank (income +, bills and transfers -)."""
    done = set((m.get("bank_check") or {}).get("done") or [])
    total = sum(v for k, v in m["income"].items() if f"income:{k}" not in done)
    total += sum(v for k, v in m["expenses"].items() if f"expenses:{k}" not in done)
    total -= sum(m[k] for k in TRANSFERS if k not in done)
    return total


def roll_forward(plan: dict) -> list[dict]:
    """The plan month by month, exactly as the workbook computes it."""
    s = plan["settings"]
    ri, rr = s["plan_investment_return"] / 12, s["plan_retirement_return"] / 12
    rows, prev = [], {k: 0.0 for k in BALANCES}
    for k, m in enumerate(plan["months"]):
        income = sum(m["income"].values())
        expenses = sum(m["expenses"].values())
        net = income + expenses
        to_bank = net - m["investment_contribution"] - m["house_fund_contribution"]
        first = k == 0
        check = m.get("bank_check")
        bank = (m["bank_balance"] if m["bank_balance"] is not None
                else check["balance"] + still_to_come(m) if check else prev["bank_balance"] + to_bank)
        bal = {
            "bank_balance": bank,
            "retirement_balance": m["retirement_balance"] if m["retirement_balance"] is not None
            else (0.0 if first else prev["retirement_balance"] * (1 + rr)) + m["retirement_contribution"],
            "investment_balance": m["investment_balance"] if m["investment_balance"] is not None
            else (0.0 if first else prev["investment_balance"] * (1 + ri)) + m["investment_contribution"],
            "house_fund_balance": m["house_fund_balance"] if m["house_fund_balance"] is not None
            else prev["house_fund_balance"] + m["house_fund_contribution"],
            "home_equity": m["home_equity"] if m["home_equity"] is not None else prev["home_equity"],
        }
        worth = sum(bal.values())
        rows.append({"month": m["month"], "income": round(income, 2), "expenses": round(expenses, 2),
                     "net_cash_flow": round(net, 2), "net_to_bank": round(to_bank, 2),
                     "investment_contribution": m["investment_contribution"],
                     "house_fund_contribution": m["house_fund_contribution"],
                     "retirement_contribution": m["retirement_contribution"],
                     **{key: round(v, 2) for key, v in bal.items()},
                     "anchored": [key for key in BALANCES if m[key] is not None],
                     "bank_check": None if not check or m["bank_balance"] is not None else
                     {"as_of": check["as_of"], "balance": check["balance"], "still_to_come": round(still_to_come(m), 2)},
                     "net_worth": round(worth, 2),
                     "net_worth_gain": None if first else round(worth - rows[-1]["net_worth"], 2)})
        prev = bal
    return rows


# ---------------------------------------------------------------- workbook import

def _month_key(value) -> str | None:
    if isinstance(value, datetime):
        return f"{value.year:04d}-{value.month:02d}"
    text = str(value or "").strip()
    match = re.match(r"^([A-Za-z]{3})-(\d{4})$", text)
    if match:
        names = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
        if match.group(1).lower() in names:
            return f"{match.group(2)}-{names.index(match.group(1).lower()) + 1:02d}"
    return MONTH.match(text) and text or None


def import_workbook(data: bytes) -> dict:
    """Turn the Income & Net Worth Tracker workbook into a plan, and check it reproduces the workbook."""
    import openpyxl

    try:
        formulas = openpyxl.load_workbook(io.BytesIO(data))
        values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable upload is the same answer
        raise HouseholdPlanError("this file is not a readable .xlsx workbook") from exc
    return import_books(formulas, values, hashlib.sha256(data).hexdigest())


def import_books(formulas, values, file_sha256: str) -> dict:
    """The import itself, from the workbook opened twice: once for formulas, once for Excel's saved values."""
    import openpyxl

    if "Monthly Tracker" not in formulas.sheetnames:
        raise HouseholdPlanError("the workbook has no 'Monthly Tracker' tab")
    f, v = formulas["Monthly Tracker"], values["Monthly Tracker"]
    header_row = next((r for r in range(1, 15) if str(v.cell(r, 1).value or "").strip().lower() == "month"), None)
    if header_row is None:
        raise HouseholdPlanError("could not find the 'Month' header on the Monthly Tracker tab")
    headers = {str(v.cell(header_row, c).value or "").strip(): c for c in range(1, v.max_column + 1) if v.cell(header_row, c).value}
    need = ["Total Income", "Total Expenses", "Investment Contribution", "House Fund Contribution", "Bank Balance",
            "Retirement Balance", "Investment Balance", "House Fund Balance", "Home Equity", "Net Worth"]
    missing = [h for h in need if h not in headers]
    if missing:
        raise HouseholdPlanError(f"the Monthly Tracker tab is missing the column {missing[0]!r}")
    col = headers
    income_lines = [h for h, c in sorted(col.items(), key=lambda x: x[1]) if c < col["Total Income"]
                    and h.lower() != "month"]
    expense_lines = [h for h, c in sorted(col.items(), key=lambda x: x[1]) if col["Total Income"] < c < col["Total Expenses"]]

    settings = dict(DEFAULT_SETTINGS)
    if "Investment Planner" in values.sheetnames:
        rate = values["Investment Planner"]["C7"].value
        if isinstance(rate, (int, float)):
            settings["plan_investment_return"] = float(rate)
    if "House Purchase Planner" in values.sheetnames:
        hp = values["House Purchase Planner"]
        for cell, key in (("C8", "down_payment_pct"), ("C9", "closing_cost_pct"), ("C19", "mortgage_rate"), ("C20", "loan_years")):
            if isinstance(hp[cell].value, (int, float)) and hp[cell].value > 0:
                settings[key] = float(hp[cell].value)

    def letter(name):
        return openpyxl.utils.get_column_letter(col[name])

    rows, workbook_rows = [], []
    rr = settings["plan_retirement_return"] / 12
    prev_ret = None
    for r in range(header_row + 1, v.max_row + 1):
        month = _month_key(v.cell(r, 1).value)
        if not month:
            continue
        cell = lambda name: v.cell(r, col[name]).value  # noqa: E731

        def anchored(name):
            formula = f.cell(r, col[name]).value
            if formula is None or formula == "":
                return False
            return not (isinstance(formula, str) and re.search(rf"\b{letter(name)}{r - 1}\b", formula))

        item = {"month": month,
                "income": {line: _number(v.cell(r, col[line]).value, line) for line in income_lines},
                "expenses": {line: _number(v.cell(r, col[line]).value, line) for line in expense_lines},
                "investment_contribution": _number(cell("Investment Contribution"), "investment"),
                "house_fund_contribution": _number(cell("House Fund Contribution"), "house fund")}
        for name, key in (("Bank Balance", "bank_balance"), ("Retirement Balance", "retirement_balance"),
                          ("Investment Balance", "investment_balance"), ("House Fund Balance", "house_fund_balance"),
                          ("Home Equity", "home_equity")):
            item[key] = _number(cell(name), key) if anchored(name) else None
        ret = _number(cell("Retirement Balance"), "retirement")
        item["retirement_contribution"] = 0.0 if item["retirement_balance"] is not None or prev_ret is None \
            else round(ret - prev_ret * (1 + rr), 2)
        prev_ret = ret
        item["note"] = ""
        rows.append(item)
        workbook_rows.append({"month": month, "net_worth": _number(cell("Net Worth"), "net worth")})
    if not rows:
        raise HouseholdPlanError("the Monthly Tracker tab has no months")
    plan = normalize_plan({"settings": settings, "income_lines": income_lines, "expense_lines": expense_lines,
                           "months": rows})
    rolled = roll_forward(plan)
    worst = max(abs(a["net_worth"] - b["net_worth"]) for a, b in zip(rolled, workbook_rows))
    plan["source"] = {"kind": "INCOME_AND_NET_WORTH_TRACKER_WORKBOOK", "file_sha256": file_sha256,
                      "imported_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
                      "workbook_end_month": workbook_rows[-1]["month"],
                      "workbook_end_net_worth": workbook_rows[-1]["net_worth"],
                      "reproduced_end_net_worth": rolled[-1]["net_worth"],
                      "largest_monthly_difference": round(worst, 2)}
    if worst > 1.0:
        raise HouseholdPlanError(f"the imported plan differs from the workbook by up to ${worst:,.2f} in a month; nothing was saved")
    return plan


# ---------------------------------------------------------------- storage (immutable versions)

@dataclass(frozen=True)
class PlanVersion:
    version_id: str
    fingerprint: str
    recorded_at: datetime
    recorded_by: str
    plan: dict

    def summary(self) -> dict:
        return {"version_id": self.version_id, "recorded_at": self.recorded_at.isoformat(), "recorded_by": self.recorded_by}


def create_version(plan: dict, recorded_by: str, recorded_at: datetime | None = None) -> PlanVersion:
    canonical = normalize_plan(plan)
    text = json.dumps(canonical, sort_keys=True, separators=(",", ":"))
    digest = hashlib.sha256(text.encode()).hexdigest()
    at = recorded_at or datetime.now(timezone.utc)
    return PlanVersion(f"household-plan-{at.strftime('%Y%m%dT%H%M%S')}-{digest[:12]}", digest, at, recorded_by, canonical)


_DDL = """CREATE TABLE IF NOT EXISTS household_plan_versions (
    version_id TEXT PRIMARY KEY,
    fingerprint TEXT NOT NULL,
    recorded_at {ts} NOT NULL,
    recorded_by TEXT NOT NULL,
    plan_json TEXT NOT NULL)"""


def _hydrate(row) -> PlanVersion:
    at = row[2] if isinstance(row[2], datetime) else datetime.fromisoformat(str(row[2]))
    return PlanVersion(str(row[0]), str(row[1]), at, str(row[3]), json.loads(row[4]))


class SQLiteHouseholdPlanRepository:
    def __init__(self, connection: sqlite3.Connection):
        self._db = connection

    def initialize(self):
        with self._db:
            self._db.execute(_DDL.format(ts="TEXT"))

    def save(self, version: PlanVersion) -> PlanVersion:
        latest = self.current()
        if latest and latest.fingerprint == version.fingerprint:
            return latest
        with self._db:
            self._db.execute("INSERT INTO household_plan_versions VALUES (?,?,?,?,?)",
                             (version.version_id, version.fingerprint, version.recorded_at.isoformat(),
                              version.recorded_by, json.dumps(version.plan, sort_keys=True)))
        return version

    def current(self) -> PlanVersion | None:
        row = self._db.execute("SELECT version_id,fingerprint,recorded_at,recorded_by,plan_json FROM household_plan_versions "
                               "ORDER BY recorded_at DESC, version_id DESC LIMIT 1").fetchone()
        return _hydrate(row) if row else None

    def history(self, limit: int = 20) -> list[PlanVersion]:
        rows = self._db.execute("SELECT version_id,fingerprint,recorded_at,recorded_by,plan_json FROM household_plan_versions "
                                "ORDER BY recorded_at DESC, version_id DESC LIMIT ?", (limit,)).fetchall()
        return [_hydrate(r) for r in rows]


class PostgresHouseholdPlanRepository:
    def __init__(self, connection_factory: Callable[[], object]):
        self._connection_factory = connection_factory

    @classmethod
    def from_dsn(cls, dsn: str):
        if not dsn.strip():
            raise ValueError("PostgreSQL DSN must not be blank")
        import psycopg
        return cls(lambda: psycopg.connect(dsn))

    def initialize(self):
        with closing(self._connection_factory()) as db, db, db.cursor() as cursor:
            cursor.execute(_DDL.format(ts="TIMESTAMPTZ"))

    def save(self, version: PlanVersion) -> PlanVersion:
        latest = self.current()
        if latest and latest.fingerprint == version.fingerprint:
            return latest
        with closing(self._connection_factory()) as db, db, db.cursor() as cursor:
            cursor.execute("INSERT INTO household_plan_versions VALUES (%s,%s,%s,%s,%s)",
                           (version.version_id, version.fingerprint, version.recorded_at, version.recorded_by,
                            json.dumps(version.plan, sort_keys=True)))
        return version

    def current(self) -> PlanVersion | None:
        with closing(self._connection_factory()) as db, db.cursor() as cursor:
            cursor.execute("SELECT version_id,fingerprint,recorded_at,recorded_by,plan_json FROM household_plan_versions "
                           "ORDER BY recorded_at DESC, version_id DESC LIMIT 1")
            row = cursor.fetchone()
        return _hydrate(row) if row else None

    def history(self, limit: int = 20) -> list[PlanVersion]:
        with closing(self._connection_factory()) as db, db.cursor() as cursor:
            cursor.execute("SELECT version_id,fingerprint,recorded_at,recorded_by,plan_json FROM household_plan_versions "
                           "ORDER BY recorded_at DESC, version_id DESC LIMIT %s", (limit,))
            rows = cursor.fetchall()
        return [_hydrate(r) for r in rows]
