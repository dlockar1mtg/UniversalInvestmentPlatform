"""Live market, economic, and official Metals providers with injectable transport."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from io import BytesIO
import json
import os
import re
from typing import Callable, Mapping, Sequence
from urllib.parse import urlencode
from urllib.request import urlopen

Transport = Callable[[str, float], bytes]


class ProviderError(RuntimeError):
    pass


def _default_transport(url: str, timeout: float) -> bytes:
    with urlopen(url, timeout=timeout) as response:  # noqa: S310 - fixed provider URLs
        return response.read()


def _request_bytes(url: str, transport: Transport, timeout: float) -> bytes:
    try:
        payload = transport(url, timeout)
    except Exception as exc:
        raise ProviderError(f"provider request failed: {type(exc).__name__}") from exc
    if not isinstance(payload, bytes) or not payload:
        raise ProviderError("provider returned an empty or invalid payload")
    return payload


def _json_request(base_url: str, parameters: Mapping[str, str], transport: Transport, timeout: float) -> dict:
    url = f"{base_url}?{urlencode(parameters)}"
    try:
        payload = json.loads(transport(url, timeout).decode("utf-8"))
    except Exception as exc:
        raise ProviderError(f"provider request failed: {type(exc).__name__}") from exc
    if not isinstance(payload, dict):
        raise ProviderError("provider returned an invalid payload")
    for key in ("Error Message", "Note", "Information", "error_message"):
        if payload.get(key):
            raise ProviderError(f"provider rejected request: {payload[key]}")
    return payload


def _decimal(value: object) -> Decimal | None:
    text = re.sub(r"[^0-9.\-]", "", str(value))
    if not text or text in {"-", ".", "-."}:
        return None
    try:
        parsed = Decimal(text)
    except InvalidOperation:
        return None
    return parsed if parsed.is_finite() else None


@dataclass(frozen=True)
class MarketQuote:
    symbol: str
    price: Decimal
    trading_date: date
    provider: str


@dataclass(frozen=True)
class EconomicObservation:
    series_id: str
    value: Decimal
    observation_date: date
    provider: str = "fred"


@dataclass(frozen=True)
class CommodityObservation:
    asset: str
    value: Decimal
    observation_date: date
    unit: str
    provider: str
    series_id: str


class AlphaVantageProvider:
    BASE_URL = "https://www.alphavantage.co/query"

    def __init__(self, api_key: str | None = None, *, transport: Transport = _default_transport, timeout: float = 15.0):
        self._api_key = (api_key or os.getenv("UIIP_ALPHA_VANTAGE_API_KEY", "")).strip()
        if not self._api_key:
            raise ValueError("UIIP_ALPHA_VANTAGE_API_KEY is required")
        self._transport = transport
        self._timeout = timeout

    def quote(self, symbol: str) -> MarketQuote:
        normalized = symbol.strip().upper()
        if not normalized:
            raise ValueError("symbol is required")
        payload = _json_request(self.BASE_URL, {
            "function": "GLOBAL_QUOTE", "symbol": normalized, "apikey": self._api_key,
        }, self._transport, self._timeout)
        quote = payload.get("Global Quote")
        if not isinstance(quote, dict) or not quote:
            raise ProviderError("Alpha Vantage returned no quote")
        try:
            price = Decimal(str(quote["05. price"]))
            trading_date = date.fromisoformat(str(quote["07. latest trading day"]))
        except (KeyError, InvalidOperation, ValueError) as exc:
            raise ProviderError("Alpha Vantage quote fields were invalid") from exc
        if not price.is_finite() or price < 0:
            raise ProviderError("Alpha Vantage quote price was invalid")
        return MarketQuote(normalized, price, trading_date, "alpha_vantage")


class FREDProvider:
    BASE_URL = "https://api.stlouisfed.org/fred/series/observations"

    def __init__(self, api_key: str | None = None, *, transport: Transport = _default_transport, timeout: float = 15.0):
        self._api_key = (api_key or os.getenv("UIIP_FRED_API_KEY", "")).strip()
        if not self._api_key:
            raise ValueError("UIIP_FRED_API_KEY is required")
        self._transport = transport
        self._timeout = timeout

    def latest(self, series_id: str) -> EconomicObservation:
        normalized = series_id.strip().upper()
        if not normalized:
            raise ValueError("series_id is required")
        payload = _json_request(self.BASE_URL, {
            "series_id": normalized, "api_key": self._api_key, "file_type": "json",
            "sort_order": "desc", "limit": "10",
        }, self._transport, self._timeout)
        observations = payload.get("observations", [])
        for observation in observations if isinstance(observations, list) else []:
            value = str(observation.get("value", "."))
            if value == ".":
                continue
            try:
                parsed = Decimal(value)
                observed = date.fromisoformat(str(observation["date"]))
            except (InvalidOperation, KeyError, ValueError) as exc:
                raise ProviderError("FRED observation fields were invalid") from exc
            if parsed.is_finite():
                return EconomicObservation(normalized, parsed, observed)
        raise ProviderError("FRED returned no numeric observation")


class _HTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "tr":
            self._row = []
        elif tag in {"td", "th"} and self._row is not None:
            self._cell = []

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"td", "th"} and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if any(self._row):
                self.rows.append(self._row)
            self._row = None


class EIAUraniumProvider:
    """Latest U.S. uranium weighted-average purchase price from EIA."""

    URL = "https://www.eia.gov/uranium/marketing/summarytable1b.php"

    def __init__(self, *, transport: Transport = _default_transport, timeout: float = 20.0):
        self._transport = transport
        self._timeout = timeout

    def latest(self) -> CommodityObservation:
        payload = _request_bytes(self.URL, self._transport, self._timeout)
        parser = _HTMLTableParser()
        try:
            document = payload.decode("utf-8", errors="replace")
            parser.feed(document)
        except Exception as exc:
            raise ProviderError("EIA uranium response was not valid HTML") from exc
        page_text = " ".join(cell for row in parser.rows for cell in row).lower()
        if "uranium" not in document.lower() or "price" not in page_text:
            raise ProviderError("EIA uranium table signature was not found")

        observations: list[tuple[int, Decimal]] = []
        for row in parser.rows:
            year_index = next(
                (index for index, cell in enumerate(row) if re.fullmatch(r"(?:19|20)\d{2}", cell.strip())),
                None,
            )
            if year_index is None:
                continue
            values = [_decimal(cell) for cell in row[year_index + 1:]]
            numeric = [value for value in values if value is not None and value >= 0]
            if numeric:
                observations.append((int(row[year_index]), numeric[-1]))
        if not observations:
            raise ProviderError("EIA uranium table contained no usable annual prices")
        year, value = max(observations, key=lambda item: item[0])
        return CommodityObservation(
            "uranium", value, date(year, 12, 31), "usd_per_pound_u3o8_equivalent",
            "eia", "EIA::URANIUM_WEIGHTED_AVG",
        )


class WorldBankCommodityProvider:
    """Latest monthly official commodity benchmarks from the World Bank Pink Sheet."""

    URL = "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx"
    TARGETS = {
        "gold": ("Gold", "usd_per_troy_ounce"),
        "silver": ("Silver", "usd_per_troy_ounce"),
        "platinum": ("Platinum", "usd_per_troy_ounce"),
        "copper": ("Copper", "usd_per_metric_ton"),
        "aluminum": ("Aluminum", "usd_per_metric_ton"),
        "nickel": ("Nickel", "usd_per_metric_ton"),
        "zinc": ("Zinc", "usd_per_metric_ton"),
        "tin": ("Tin", "usd_per_metric_ton"),
    }

    def __init__(self, *, transport: Transport = _default_transport, timeout: float = 30.0):
        self._transport = transport
        self._timeout = timeout

    @staticmethod
    def _month(value: object) -> date | None:
        if isinstance(value, datetime):
            return date(value.year, value.month, 1)
        if isinstance(value, date):
            return date(value.year, value.month, 1)
        text = str(value).strip()
        match = re.fullmatch(r"(\d{4})M(\d{1,2})", text, re.IGNORECASE)
        if match:
            year, month = map(int, match.groups())
            try:
                return date(year, month, 1)
            except ValueError:
                return None
        try:
            parsed = date.fromisoformat(text[:10])
        except ValueError:
            return None
        return date(parsed.year, parsed.month, 1)

    def latest(self, assets: Sequence[str] | None = None) -> tuple[CommodityObservation, ...]:
        requested = tuple(dict.fromkeys(asset.lower() for asset in (assets or self.TARGETS)))
        unknown = sorted(set(requested) - set(self.TARGETS))
        if unknown:
            raise ValueError(f"unsupported World Bank commodities: {', '.join(unknown)}")
        payload = _request_bytes(self.URL, self._transport, self._timeout)
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        except Exception as exc:
            raise ProviderError("World Bank response was not a valid XLSX workbook") from exc

        sheet = next((item for item in workbook.worksheets if "monthly" in item.title.lower()), workbook.worksheets[0])
        rows = list(sheet.iter_rows(values_only=True))
        header_index = next(
            (
                index for index, row in enumerate(rows[:30])
                if "gold" in {str(value).strip().lower() for value in row if value is not None}
                and any(str(value).strip().lower().startswith(("copper", "silver")) for value in row if value is not None)
            ),
            None,
        )
        if header_index is None:
            raise ProviderError("World Bank monthly commodity header was not found")
        headers = [str(value or "").strip() for value in rows[header_index]]
        columns: dict[str, int] = {}
        for asset in requested:
            label = self.TARGETS[asset][0].lower()
            column = next((index for index, value in enumerate(headers) if value.lower().startswith(label)), None)
            if column is None:
                raise ProviderError(f"World Bank workbook is missing the {asset} column")
            columns[asset] = column

        latest: dict[str, tuple[date, Decimal]] = {}
        for row in rows[header_index + 1:]:
            if not row:
                continue
            observed = self._month(row[0])
            if observed is None:
                continue
            for asset, column in columns.items():
                if column >= len(row):
                    continue
                value = _decimal(row[column])
                if value is not None and value >= 0:
                    current = latest.get(asset)
                    if current is None or observed > current[0]:
                        latest[asset] = (observed, value)
        missing = [asset for asset in requested if asset not in latest]
        if missing:
            raise ProviderError(f"World Bank workbook has no usable observations for: {', '.join(missing)}")
        return tuple(
            CommodityObservation(
                asset, latest[asset][1], latest[asset][0], self.TARGETS[asset][1],
                "world_bank", f"WORLD_BANK::{asset.upper()}_MONTHLY",
            )
            for asset in requested
        )
