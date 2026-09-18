from __future__ import annotations

import argparse
import csv
from io import BytesIO
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from foundation.production.providers import WorldBankCommodityProvider, _decimal, _request_bytes


def collect_rows() -> list[dict[str, object]]:
    provider = WorldBankCommodityProvider()
    workbook_url = provider._resolve_workbook_url()
    payload = _request_bytes(workbook_url, provider._transport, provider._timeout)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise RuntimeError("World Bank response was not a valid XLSX workbook") from exc

    sheet = next((item for item in workbook.worksheets if "monthly" in item.title.lower()), workbook.worksheets[0])
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (
            index for index, row in enumerate(rows[:30])
            if "gold" in {str(value).strip().lower() for value in row if value is not None}
            and any(str(value).strip().lower().startswith(("copper", "silver")) for value in row if value is not None)
        ),
        None,
    )
    if header_index is None:
        raise RuntimeError("World Bank monthly commodity header was not found")

    headers = [str(value or "").strip() for value in rows[header_index]]
    columns: dict[str, int] = {}
    for asset, (label, _) in provider.TARGETS.items():
        column = next((i for i, value in enumerate(headers) if value.lower().startswith(label.lower())), None)
        if column is None:
            raise RuntimeError(f"World Bank workbook is missing the {asset} column")
        columns[asset] = column

    output: list[dict[str, object]] = []
    for row in rows[header_index + 1:]:
        if not row:
            continue
        observed = provider._month(row[0])
        if observed is None:
            continue
        for asset, column in columns.items():
            if column >= len(row):
                continue
            value = _decimal(row[column])
            if value is None or value < 0:
                continue
            output.append(
                {
                    "asset_id": asset,
                    "observation_date": observed.isoformat(),
                    "value": str(value),
                    "source": "world_bank",
                    "unit": provider.TARGETS[asset][1],
                    "series_id": f"WORLD_BANK::{asset.upper()}_MONTHLY",
                }
            )
    if not output:
        raise RuntimeError("World Bank monthly history returned no usable rows")
    return sorted(output, key=lambda row: (str(row["asset_id"]), str(row["observation_date"])))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    rows = collect_rows()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["asset_id", "observation_date", "value", "source", "unit", "series_id"])
        writer.writeheader()
        writer.writerows(rows)
    print(f"METALS BENCHMARK HISTORY COLLECTION: PASS ({len(rows)} rows)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
